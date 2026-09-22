from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
footer=sec.footer.paragraphs[0]
footer.alignment=2
footer.add_run('MongoDB Atlas  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
def sub(t): doc.add_heading(t,2)
def bullets(items, check=False):
    for t in items:
        doc.add_paragraph(('☐ '+t) if check else t, style=None if check else 'List Bullet')
def page(): doc.add_page_break()
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,s in zip(t.rows[0].cells,headers): c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row): c.text=s
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trpr.append(no)
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF2' if i==0 else ('F7F7F7' if i%2==0 else 'FFFFFF')); pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(5); pp.paragraph_format.space_before=Pt(5)
                for r in pp.runs: r.font.size=Pt(10); r.bold=(i==0)
    doc.add_paragraph()

doc.add_paragraph('Manual de onboarding de MongoDB Atlas', 'Title')
p('Aplicaciones alojadas en Scotia Atlas Platform')
p('Propuesta operativa con requisitos, responsables, solicitudes y validaciones. Completar los campos entre corchetes antes de abrir los tickets.')
h('1 Objetivo y alcance')
p('Habilitar MongoDB Atlas para la aplicación, incluyendo aprobaciones, presupuesto, clúster, cifrado, conectividad privada, credenciales, monitoreo, logs y alertas. El equipo de aplicación coordina los tickets y valida la conexión.')
p('MongoDB Atlas es el servicio de base de datos. Scotia Atlas Platform es la plataforma del banco donde se aloja la aplicación. El procedimiento aplica a IST, UAT, NFT y PROD; NON-IST comprende UAT, NFT y PROD.')
h('2 Ficha de la aplicación')
table(['Dato','Información por completar'],[
('Aplicación y nombre corto','[Nombre] / [Nombre corto]'),('Responsable y contacto','[Equipo / contacto]'),('EPM y CIAD code','[EPM] / [CIAD]'),('Financiamiento','[OP Code] / [Transit#]'),('Ambiente y región','[IST / UAT / NFT / PROD] / [Región]'),('Scotia Atlas Platform','[Clúster / referencia de onboarding]'),('Seguridad','[Security Advisor] / [Clasificación de datos]'),('Grupo AD y fecha requerida','[Grupo] / [Fecha]')],[2.2,4.4])
h('3 Requisitos obligatorios')
bullets(['Aprobación GTEP ARB obtenida.','Diseño inicial y presupuesto del clúster completados.','Financiamiento aprobado con OP Code, Transit# y CIAD code identificado.','TRA iniciado y Security Advisor asignado.','Onboarding de la aplicación a Scotia Atlas Platform iniciado.'],True)

page(); h('4 Presupuesto y coordinación')
table(['Concepto','Referencia presupuestaria'],[('Clúster MongoDB Atlas','Cotización según la carga de la aplicación.'),('Infraestructura interna PSC','Aproximadamente $2.000 mensuales.'),('Licencia Dynatrace','$6.500 por única vez.'),('Costos de equipos internos','Consultar la hoja de costos interna.')],[2.2,4.4])
p('Confirmar moneda, vigencia y alcance de estas cifras antes de aprobar el presupuesto. Contacto indicado para estimar el clúster: Abbhilash Mahesh, abbhi.mahesh@mongodb.com.')
p('El equipo de aplicación debe acordar con cada equipo los plazos de onboarding por ambiente. Como preparación adicional, revisar con DBA capacidad, carga esperada, crecimiento, disponibilidad, respaldo y recuperación; estos puntos son una propuesta de preparación y no una plantilla oficial del documento base.')
h('5 Responsables por ambiente')
table(['Equipo','Responsabilidad'],[
('Aplicación','Tickets, aprobaciones, coordinación, FPR y validación funcional.'),('DBA','Clúster y configuración; integración del lado de MongoDB. Credenciales y API keys en IST.'),('COPS','Cuenta de servicio, clave JSON y VPC Service Controls. También keyring y cifrado en IST.'),('GNE y Redes','PSC del lado de GCP, datos de conectividad y ejecución del firewall.'),('Cloud Crypto','Keyring y clave de cifrado en NON-IST; ruta de Vault y coordinación de secretos.'),('GIAM','Credenciales y API keys en NON-IST; coordinación de secretos y configuración de cifrado en MongoDB. Gestión de reset de cuentas FC en producción.'),('PCM / ESLM / GEMS','Dynatrace / logs / notificaciones de alertas.')],[1.5,5.1])
p('Secuencia de coordinación propuesta: requisitos y DOU → solicitudes DBA/KMS → configuración PSC → FPR → conexión con credenciales de Vault → validaciones. Confirmar qué actividades pueden ejecutarse en paralelo y sus dependencias.')

page(); h('6 DOU y aprobaciones de seguridad')
p('Responsable de gestión: equipo de aplicación. Preparar el Documento de Entendimiento (DOU) con la plantilla interna GCP Non-Atlas e iniciar la consulta tecnológica y los cambios requeridos por COPS.')
bullets(['Incluir aplicación, ambiente, diseño, integración con KMS, cuenta de servicio, PSC y tickets relacionados.','Adjuntar aprobación del Security Advisor para la rotación de la clave JSON y la clave de cifrado, según la clasificación de datos.','Adjuntar aprobación para actualizar VPC Service Controls con las IP públicas de MongoDB y la cuenta de servicio.','DBA entrega las IP públicas después del despliegue; su incorporación puede necesitar un cambio separado.','Reutilizar el DOU de IST para NON-IST, incorporando los requisitos correspondientes.'])
p('Referencias DOU: DMREQ0023678 / DMREQ0024433. Resultado esperado: documento y aprobaciones disponibles para los cambios.')
h('7 Solicitud de despliegue a DBA')
p('Grupo: GTS - DBO - MONGODB. Asunto: Onboarding MongoDB Atlas — [Aplicación] — [Ambiente]. Adjuntar ficha, diseño, aprobaciones y DOU.')
bullets(['Crear la estructura de organización/proyecto que corresponda y desplegar el clúster según el diseño.','Aplicar etiquetas, incluido CIAD code, y comunicar la información de facturación.','Configurar el lado MongoDB para KMS y PSC y apoyar monitoreo, logs y alertas.','En IST, crear las credenciales; en NON-IST, coordinar con GIAM.','Entregar insumos para tickets dependientes, incluido el script de PSC cuando corresponda.'])
p('Referencia: REQ5436024 / RITM6329001 / TASK7144344. Resultado: clúster y datos técnicos disponibles para las integraciones.')
h('8 Cifrado con GCP KMS')
sub('IST')
p('COPS crea keyring, clave de cifrado, cuenta de servicio y clave JSON; actualiza VPC Service Controls y entrega la clave JSON y Key Version ID a DBA mediante el mecanismo autorizado. DBA configura y valida KMS desde MongoDB.')
sub('UAT NFT y PROD')
p('Cloud Crypto crea keyring y clave de cifrado con la rotación aprobada. COPS crea cuenta de servicio y JSON, entrega el JSON a PAM y actualiza VPC Service Controls. Coordinar la configuración del cifrado con GIAM, DBA, COPS y Cloud Crypto.')
p('Grupo COPS: GTS - CODO - COPS – Public. Proyectos indicados: nbyzd-0136-mongodb (no productivo) y pbyzd-0136-mongod (productivo). Confirmar identificadores. Referencias CKI: CKI-12871 / CKI-13570. Resultado: integración KMS validada.')

page(); h('9 Conectividad privada mediante PSC')
p('Responsables: GNE y DBA. El equipo de aplicación incorpora PSC al DOU, abre el cambio y lo vincula con DBA. En NON-IST debe incluir el script generado por DBA desde MongoDB.')
bullets(['GNE configura PSC del lado GCP. DBA configura el lado MongoDB y valida la integración.','En IST, el documento describe la asignación a GNE desde el grupo COPS.','En NON-IST, usar la plantilla de cambio correspondiente a no productivo o producción.','Solicitar a GNE el archivo JSON con las IP privadas de endpoints PSC para el FPR.'])
p('Grupos: GTS - CODO - COPS – Public y GTEP - GNE - BUILD Data Center and Cloud. Resultado: PSC configurado e IP privadas disponibles.')
h('10 Apertura de firewall mediante FPR')
p('Portal: https://algosec.security.bns/. Crear una solicitud FireFlow con la plantilla “Scotiabank-GCP template”. Incluir [gcp-psc] en el asunto y los comentarios.')
p('Asunto propuesto: [gcp-psc] Conectividad de [Aplicación] hacia MongoDB Atlas — [Ambiente].')
bullets(['Orígenes: subredes y etiquetas vigentes de la aplicación y jumpboxes autorizados.','Destinos: IP privadas de los endpoints PSC entregadas por GNE.','Puertos listados en el documento: 27017 (MongoDB), 27016 (mongos) y 27015 (BI Connector). Confirmar con DBA/GNE cuáles aplican.','Validar conexión desde la aplicación y los jumpboxes autorizados después del cambio.'])
p('En producción, no incluir jumpboxes salvo necesidad de acceso de solo lectura para soporte y previa consulta con el Security Advisor. Verificar subredes actuales: el documento advierte una migración de tenant en no productivo y los ejemplos pueden estar desactualizados.')
h('11 Usuarios y secretos en Vault')
p('IST: DBA crea el usuario y contraseña; el equipo de aplicación carga los secretos en HashiCorp Vault. NON-IST: GIAM crea credenciales y trabaja con Cloud Crypto para la ruta y carga en Vault. Vincular la solicitud GIAM con CKI.')
p('Grupo GIAM: Database ID Provisioning. Referencia CKI para Vault: CKI-13761. Confirmar quién presenta este ticket, pues el documento asigna esa actividad tanto a GIAM como a la aplicación.')
bullets(['IST: app_<EPM-code>_<app_short_name>_i_1','UAT: app_<EPM-code>_<app_short_name>_u_1','NFT: app_<EPM-code>_<app_short_name>_n_1','PROD: app_<EPM-code>_<app_short_name>_p_1'])
p('En todos los ambientes, la aplicación consume las credenciales desde Vault mediante Kubernetes External Secrets. Los usuarios personales se gestionan separadamente; en IST coordinar con DBA su nomenclatura y creación.')

page(); h('12 Monitoreo logs y alertas')
p('Dynatrace, ESLM y GEMS son obligatorios en UAT, NFT y PROD; en IST son opcionales según la decisión del equipo de aplicación. El intake PCM de MongoDB es independiente del de Scotia Atlas Platform.')
table(['Integración','Solicitud y grupo'],[('Dynatrace / PCM','Presentar EPM, nombre de aplicación y grupo AD. Grupo: GTS - Tools & Monitoring - Dynatrace Support. Referencia CTASK0104193.'),('ESLM','Solicitar integración de logs. Grupo: IS&C – CSS - CIA. Referencia RITM5913831.'),('GEMS','Solicitar notificaciones de alertas. Grupo: GTS - Tools & Monitoring - GEMS. Referencia GEMS0009591.')],[1.5,5.1])
p('API keys: DBA las genera en IST; GIAM (Key Platform Provisioning), en NON-IST. Coordinar su entrega a monitoreo/logs mediante el proceso autorizado. Resultado: métricas, registros y alertas disponibles.')
h('13 Acceso a jumpbox cuando corresponda')
bullets(['Solicitar acceso en https://lam.bns/iiq/login.jsf.','Rol/grupo indicado: APP-BJ8H-MongoDBUsers.','Confirmar el jumpbox autorizado y coordinar credenciales personales de base de datos.','Las herramientas cliente están instaladas en los jumpboxes COPS según el documento.','Mantener las restricciones de producción indicadas en el apartado FPR.'])
h('14 Validación y cierre propuestos')
bullets(['Clúster en el ambiente correcto y conforme al diseño.','CIAD code y etiquetas aplicados.','Integración KMS validada.','PSC configurado y FPR implementado.','Usuario creado y secretos disponibles en Vault.','Kubernetes External Secrets configurado.','Conexión y autenticación desde la aplicación verificadas.','Operación funcional de prueba validada según permisos de la aplicación.','Dynatrace, ESLM y GEMS validados cuando correspondan.','Accesos de soporte habilitados cuando correspondan.','Evidencias y referencias de tickets registradas.'],True)
p('La prueba funcional y esta lista de cierre son una propuesta para verificar la habilitación con los equipos responsables.')

page(); h('15 Matriz de seguimiento')
table(['Frente','Ambiente','Ticket y estado'],[('Consulta tecnológica y DOU','Todos / reutilizar DOU','[Completar]'),('Despliegue DBA','Todos','[Completar]'),('Cuenta de servicio y VPC SC con COPS','Todos','[Completar]'),('Keyring y cifrado con COPS','IST','[Completar]'),('Keyring y cifrado con Cloud Crypto','NON-IST','[Completar]'),('PSC con GNE y DBA','Todos','[Completar]'),('Firewall FPR','Todos','[Completar]'),('Usuario y Vault con DBA/aplicación','IST','[Completar]'),('Credenciales con GIAM','NON-IST','[Completar]'),('Vault con Cloud Crypto / CKI','NON-IST','[Completar]'),('PCM / ESLM / GEMS','NON-IST obligatorio','[Completar]'),('Acceso a jumpbox','Según necesidad','[Completar]')],[3.0,1.8,1.8])
p('La matriz identifica frentes de trabajo; no implica necesariamente un ticket por fila. Confirmar solicitudes, cambios y tareas requeridos por cada grupo.')
h('16 Consultas pendientes')
bullets(['¿Cuál es la plantilla vigente para cada ticket y ambiente?','¿Qué dimensionamiento y configuración deben adjuntarse a DBA?','¿Cuáles son las subredes, etiquetas y puertos aplicables?','¿Quién presenta el ticket CKI de Vault: aplicación o GIAM?','¿Quién recibe y configura cada dato KMS en NON-IST?','¿Siguen vigentes los costos y en qué moneda están expresados?','¿Cuáles son los plazos y ventanas de cambio?','¿Cuáles son los enlaces actuales de los catálogos ESLM y GEMS?'])
h('17 Referencia y puntos por confirmar')
p('Fuente: MongoDB Atlas Onboarding Process, creado por Balraj Singh y actualizado por Mustafiz Rahman el 13 de agosto de 2026. El documento base está marcado IN PROGRESS (DRAFT). Solicitar las plantillas, diagramas y hoja de costos asociados.')
p('Confirmar tres inconsistencias del original: una mención a NON-IST dentro del procedimiento IST para secretos; la responsabilidad de abrir CKI para Vault; y la entrega de Key Version ID a PAM, que aparece con una duda explícita. Las secciones detalladas asignan la carga en Vault a la aplicación en IST y a GIAM/Cloud Crypto en NON-IST.')

page(); h('Anexo Propuesta de ticket inicial a DBA')
p('Asunto: Solicitud de onboarding MongoDB Atlas — [Aplicación] — [Ambiente]')
p('Hola, equipo:')
p('Solicitamos iniciar el onboarding de MongoDB Atlas para [nombre de la aplicación], alojada o en proceso de onboarding en Scotia Atlas Platform, para el ambiente [ambiente].')
sub('Datos de la solicitud')
bullets(['Aplicación y nombre corto: [Completar].','EPM / CIAD code: [Completar].','OP Code / Transit#: [Completar].','Responsable y contacto: [Completar].','Región y fecha requerida: [Completar].'])
sub('Estado de los prerrequisitos')
bullets(['GTEP ARB: [Aprobado/Pendiente y referencia].','Diseño y presupuesto: [Adjunto/Pendiente].','Financiamiento: [Aprobado/Pendiente].','TRA y Security Advisor: [Referencia y nombre].','Onboarding Scotia Atlas: [Estado y referencia].','DOU y aprobaciones de seguridad: [Referencias/Pendiente].'])
sub('Alcance solicitado')
p('Crear y configurar el proyecto y clúster conforme al diseño acordado; aplicar las etiquetas de facturación; coordinar el lado MongoDB para KMS y PSC; entregar insumos para tickets dependientes; y apoyar credenciales, monitoreo, logs y alertas según el ambiente.')
p('Nuestro equipo gestionará los tickets de COPS, GNE, FPR y demás equipos involucrados, y validará la conexión desde la aplicación.')
sub('Consultas para iniciar')
p('Solicitamos confirmar si la información es suficiente, la plantilla vigente, el dimensionamiento requerido, las dependencias previas y el plazo estimado de atención. Por favor indicar los insumos que entregará DBA para PSC, KMS y credenciales.')
sub('Criterio de cierre propuesto')
p('Clúster creado y configurado, etiquetas aplicadas y actividades del lado MongoDB incluidas en este ticket completadas, con evidencias y referencias a las solicitudes dependientes.')
p('Adjuntos: [Ficha de aplicación], [Diseño], [Aprobaciones], [DOU] y [Tickets relacionados].')
p('Gracias.\n[Nombre y equipo]')

doc.core_properties.title='Manual de onboarding de MongoDB Atlas'
doc.core_properties.subject='Requisitos, responsables, tickets y validaciones'
doc.save(OUT/'Manual_onboarding_MongoDB_Atlas.docx')
print(OUT/'Manual_onboarding_MongoDB_Atlas.docx')
