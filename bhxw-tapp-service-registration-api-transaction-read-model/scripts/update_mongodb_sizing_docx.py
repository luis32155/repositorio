#!/usr/bin/env python3
"""Update the completed MongoDB sizing questionnaire with 2027-2032 data."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "mongodb-sizing" / "MongoDB_Sizing_Questionnaire_TAPP_Completado.docx"
OUTPUT = ROOT / "outputs" / "mongodb-sizing" / "MongoDB_Sizing_Questionnaire_TAPP_Actualizado_2027-2032.docx"
VOLUME_IMAGE = ROOT / "docs" / "cdc-workaround" / "volumetria-cce-2027-2032.png"

NAVY = "1F3A5F"
BLUE = "D9EAF7"
LIGHT_BLUE = "EFF6FB"
RED = "CC092F"
LIGHT_RED = "FCE8EC"
GREEN = "207A4B"
LIGHT_GREEN = "E7F4EC"
GRAY = "5B6573"
LIGHT_GRAY = "F2F4F6"
WHITE = "FFFFFF"


ANSWERS = {
    1: (
        "Con retención de 365 días y una proyección compacta de 3 KB: 60,9 GB "
        "(conservador) y 75,0 GB (optimista) para 2032. Con 20% para índices y "
        "30% de holgura: 95 GB y 117 GB por nodo, respectivamente. Capacidad "
        "operativa sugerida para Fase 3: 150 GB por nodo, sin incluir backups ni oplog."
    ),
    2: (
        "Usar 3 KB como supuesto provisional para el documento BSON compacto. "
        "No usar 3 KB como máximo XML: ReqPay midió 2,18 KiB realista, 5,34 KiB "
        "conservador y 9,49 KiB en la muestra máxima, todos sin firma."
    ),
    3: (
        "Usar 3 KB provisionalmente cuando la lectura devuelve la proyección compacta. "
        "Si la API retorna XML, dimensionar por endpoint y percentiles; RespListAccount "
        "crece con el número de cuentas y supera 3 KiB desde siete cuentas sin firma."
    ),
    4: (
        "Supuesto provisional: pico sostenido de 2 horas. La proyección solo entrega "
        "volumen mensual/anual; la duración y distribución horaria deben confirmarse "
        "con Negocio o mediante métricas productivas."
    ),
    5: (
        "2032 conservador: 25,75 escrituras/s pico. 2032 optimista: 31,71 escrituras/s "
        "pico. Cálculo con cuatro eventos por transacción y factor pico 10x. Prueba "
        "objetivo: 40 escrituras/s; estrés: 80 escrituras/s."
    ),
    6: (
        "Supuesto provisional: 2 horas diarias, alineadas al pico transaccional. "
        "Sustituir por la ventana real cuando se disponga de telemetría."
    ),
    7: (
        "Con tres consultas por transacción y el pico optimista 2032: 23,78 lecturas/s. "
        "Prueba objetivo: 30 lecturas/s; estrés: 60 lecturas/s. Resultado habitual: "
        "un documento por consulta de identidad."
    ),
    8: (
        "Reservar inicialmente 20% del tamaño de datos: hasta 12,18 GB en el escenario "
        "conservador y 15 GB en el optimista para 2032. Validar con collStats y "
        "$indexStats usando datos BSON representativos."
    ),
    9: "Hot data propuesto: 30 días. Retención total: 365 días mediante índice TTL.",
    10: "95% búsquedas simples/CRUD y 5% agregaciones o analítica.",
    11: (
        "Promedio: un documento MongoDB por consulta. La respuesta funcional "
        "ListAccount puede contener múltiples cuentas dentro del mismo mensaje XML."
    ),
    12: (
        "Supuesto del cuestionario: 20 minutos. Para este microservicio backend son "
        "más relevantes el TPS, la concurrencia y las consultas por transacción."
    ),
    13: (
        "Carga continua desde Kafka mediante upsert idempotente. Volumen anual: "
        "1,1-20,3 MM (conservador) y 2-25 MM (optimista) entre 2027 y 2032. "
        "Con cuatro eventos, estos aumentan las operaciones de actualización, no el "
        "número final de documentos consolidados."
    ),
    14: (
        "Sí, propuesto y pendiente de aprobación. Proyección Kafka-MongoDB: "
        "p95 <= 500 ms y p99 <= 1 000 ms en operación normal."
    ),
    15: (
        "Objetivos propuestos: lectura MongoDB/API p95 <= 100 ms y p99 <= 200 ms; "
        "upsert MongoDB p95 <= 100 ms; proyección extremo a extremo p95 <= 500 ms."
    ),
    16: (
        "Consultas estáticas definidas en el código. No se prevén consultas ad-hoc "
        "de usuario sobre el read model transaccional."
    ),
    17: (
        "No existe batch recurrente; la ingesta es continua. Para carga inicial o "
        "backfill se propone 00:00-04:00 (hora de Lima), con throttling y monitoreo del lag."
    ),
    18: (
        "Purga automática mediante TTL a los 365 días desde completed_at. Si existe "
        "retención regulatoria adicional, archivar antes del vencimiento."
    ),
    19: (
        "Producción: backup continuo/PITR más snapshot diario. RPO propuesto <= 5 min "
        "y RTO propuesto <= 60 min; pendiente de la política corporativa y validación oficial."
    ),
    20: "DEV, QA/SIT, UAT/Staging y Performance. DR forma parte de continuidad productiva.",
    21: (
        "DEV: 10%; QA/SIT: 25%; UAT/Staging: 50%; Performance: 100% temporal durante "
        "campañas. Aplicar los porcentajes sobre la capacidad de la fase correspondiente."
    ),
}


def set_font(run, size: float = 9.0, bold: bool = False, color: str = "222222") -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def clear_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05


def put_cell_text(cell, text: str, *, bold: bool = False, color: str = "222222", size: float = 8.5) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header_fill: str = NAVY, first_col_fill: str | None = LIGHT_GRAY) -> None:
    for column_index, cell in enumerate(table.rows[0].cells):
        shade_cell(cell, header_fill)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_font(run, size=8.5, bold=True, color=WHITE)
    for row in table.rows[1:]:
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if first_col_fill and column_index == 0:
                shade_cell(cell, first_col_fill)
            else:
                shade_cell(cell, WHITE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.highlight_color = None
                    r_pr = run._element.get_or_add_rPr()
                    run_shading = r_pr.find(qn("w:shd"))
                    if run_shading is not None:
                        r_pr.remove(run_shading)
                    set_font(run, size=8.5, bold=(column_index == 0), color="222222")


def add_heading(doc, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.left_indent = Twips(0)
    paragraph.paragraph_format.right_indent = Twips(0)
    paragraph.paragraph_format.first_line_indent = Twips(0)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    p_pr = paragraph._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level - 1))
    p_pr.append(outline)
    run = paragraph.add_run(text)
    set_font(run, size=15 if level == 1 else 11, bold=True, color=NAVY if level == 1 else RED)


def add_note(doc, label: str, text: str, fill: str, accent: str, *, spacer: bool = True) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9240])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(f"{label}: ")
    set_font(label_run, size=9.5, bold=True, color=accent)
    text_run = paragraph.add_run(text)
    set_font(text_run, size=9.5, color="222222")
    if spacer:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_simple_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = doc.styles.get_by_id("TableGrid", WD_STYLE_TYPE.TABLE)
    for index, header in enumerate(headers):
        put_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=8.0)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            put_cell_text(cells[index], value, size=8.0)
            if index > 0 and len(value) < 30:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths)
    style_table(table, header_fill=NAVY, first_col_fill=LIGHT_BLUE)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def update_document() -> None:
    doc = Document(SOURCE)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    doc.core_properties.title = "MongoDB Sizing Questionnaire - TAPP - Actualizado 2027-2032"
    doc.core_properties.subject = "Sizing MongoDB, volumetría CCE y prueba de tamaños XML"
    doc.core_properties.author = "TAPP Architecture"
    doc.core_properties.last_modified_by = "TAPP Architecture"

    title = doc.paragraphs[0]
    title.text = "MongoDB Sizing Questionnaire"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    for run in title.runs:
        set_font(run, size=20, bold=True, color=NAVY)

    subtitle = doc.paragraphs[1]
    subtitle.text = "TAPP | Actualización CCE 2027-2032"
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(2)
    for run in subtitle.runs:
        set_font(run, size=11, bold=True, color=RED)

    metadata = doc.paragraphs[2]
    metadata.text = "Fecha de actualización: 1 de septiembre de 2026 | Estado: estimación para validación"
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(9)
    for run in metadata.runs:
        set_font(run, size=8.5, color=GRAY)

    questionnaire = doc.tables[0]
    put_cell_text(
        questionnaire.rows[0].cells[0],
        "Cuestionario: preguntas y respuestas actualizadas",
        bold=True,
        color=WHITE,
        size=9,
    )
    for row_index, answer in ANSWERS.items():
        put_cell_text(questionnaire.rows[row_index].cells[1], answer, size=8.3)
    set_table_geometry(questionnaire, [4140, 5100])
    style_table(questionnaire, header_fill=NAVY, first_col_fill=LIGHT_GRAY)
    questionnaire.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    doc.add_page_break()
    add_heading(doc, "Anexo 1. Resultado de la validación de tamaños XML", 1)
    add_note(
        doc,
        "Conclusión",
        "3 KB puede representar un mensaje simple sin firma, pero no es un máximo válido para los XML. "
        "La muestra ReqPay poblada alcanzó 9,49 KiB sin firma y 11,71 KiB con firma dummy.",
        LIGHT_RED,
        RED,
    )
    add_heading(doc, "ReqPay", 2)
    add_simple_table(
        doc,
        ["Perfil", "Sin firma", "Con firma dummy", "Evaluación de 3 KiB"],
        [
            ["Realista", "2 227 B / 2,18 KiB", "4 498 B / 4,39 KiB", "Solo unsigned cabe"],
            ["Conservador", "5 467 B / 5,34 KiB", "7 738 B / 7,56 KiB", "No cabe"],
            ["Máximo recibido", "9 721 B / 9,49 KiB", "11 992 B / 11,71 KiB", "No cabe"],
        ],
        [2200, 1800, 1900, 3340],
    )
    add_heading(doc, "RespListAccount", 2)
    add_simple_table(
        doc,
        ["Cuentas", "Sin firma", "Con firma dummy", "Observación"],
        [
            ["1", "901 B / 0,88 KiB", "3 172 B / 3,10 KiB", "La firma supera 3 KiB"],
            ["5 (muestra)", "2 363 B / 2,31 KiB", "4 634 B / 4,53 KiB", "Evidencia original"],
            ["7", "3 397 B / 3,32 KiB", "5 668 B / 5,54 KiB", "Primer unsigned > 3 KiB"],
            ["10", "4 645 B / 4,54 KiB", "6 916 B / 6,75 KiB", "Crecimiento lineal"],
            ["20", "8 805 B / 8,60 KiB", "11 076 B / 10,82 KiB", "Requiere mayor envolvente"],
            ["50", "21 285 B / 20,79 KiB", "23 556 B / 23,00 KiB", "32 KiB todavía lo cubre"],
        ],
        [1400, 1900, 2000, 3940],
    )
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(6)
    run = note.add_run(
        "Nota metodológica: UTF-8 sin BOM y sin compresión. La firma XMLDSig agrega 2 271 bytes; "
        "sus valores son dummy y no tienen validez criptográfica. Todos los fixtures son XML bien formado."
    )
    set_font(run, size=8.5, color=GRAY)

    doc.add_page_break()
    add_heading(doc, "Anexo 2. Proyección volumétrica y capacidad", 1)
    if VOLUME_IMAGE.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(VOLUME_IMAGE), width=Inches(5.30))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(7)
        run = caption.add_run("Figura 1. Proyección volumétrica CCE revisada con el equipo.")
        set_font(run, size=8, color=GRAY)

    add_heading(doc, "Throughput de diseño", 2)
    add_simple_table(
        doc,
        ["Escenario 2032", "Transacciones/año", "TPS pico 10x", "Escrituras/s pico (4 eventos)"],
        [
            ["Conservador", "20,3 MM", "6,44", "25,75"],
            ["Optimista", "25,0 MM", "7,93", "31,71"],
        ],
        [2500, 2000, 1800, 2940],
    )
    add_note(
        doc,
        "Prueba recomendada",
        "40 escrituras/s y 30 lecturas/s como objetivo mínimo; 80 escrituras/s y 60 lecturas/s para estrés.",
        LIGHT_GREEN,
        GREEN,
    )

    add_heading(doc, "Almacenamiento por fase", 2)
    add_simple_table(
        doc,
        ["Fase", "Horizonte", "Optimista: datos a 3 KB", "Con índices + holgura", "Capacidad sugerida/nodo"],
        [
            ["1", "2027-2028", "18,0 GB", "28,08 GB", "40 GB"],
            ["2", "2029-2030", "45,0 GB", "70,20 GB", "100 GB"],
            ["3", "2031-2032", "75,0 GB", "117,00 GB", "150 GB"],
        ],
        [900, 1500, 1900, 2000, 2940],
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(
        "Supuestos: un documento consolidado por transacción, 3 000 bytes/documento, retención 365 días, "
        "20% para índices y 30% de holgura. Oplog, backups y PITR no están incluidos."
    )
    set_font(run, size=8.5, color=GRAY)

    add_heading(doc, "Decisiones pendientes", 2)
    add_note(
        doc,
        "Por confirmar",
        "si MongoDB guarda XML crudo; la distribución BSON p50/p95/p99; eventos por transacción; "
        "factor y duración del pico; cuentas promedio/p95/máximo en ListAccount; y la aprobación "
        "oficial de retención, SLA, RPO y RTO.",
        LIGHT_GRAY,
        NAVY,
        spacer=False,
    )

    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    tail_run = tail.add_run(" ")
    set_font(tail_run, size=1, color=WHITE)
    tail_run.font.hidden = True

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.text = "TAPP | MongoDB Sizing | Actualización 2027-2032"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        set_font(run, size=8, color=GRAY)

    section_type = doc._element.body.sectPr.find(qn("w:type"))
    if section_type is not None:
        section_type.set(qn("w:val"), "continuous")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    update_document()
