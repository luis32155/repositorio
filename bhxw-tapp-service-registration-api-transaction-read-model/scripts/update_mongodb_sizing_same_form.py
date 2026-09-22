#!/usr/bin/env python3
"""Update only answer cells while preserving the original Word form."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "mongodb-sizing" / "MongoDB_Sizing_Questionnaire_TAPP_Completado.docx"
OUTPUT = ROOT / "outputs" / "mongodb-sizing" / "MongoDB_Sizing_Questionnaire_TAPP_Actualizado_2027-2032_Mismo_Formato.docx"
EXPECTED_SOURCE_SHA256 = "A08FB95FBFE667C331A289D9BA1B71A577035B9080B46F9F607A17F07C4D0EF9"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
W = f"{{{W_NS}}}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


ANSWERS: dict[int, list[str]] = {
    1: [
        "2032: 60,9 GB conservador y 75,0 GB optimista sin comprimir.",
        "Con 20% de índices y 30% de holgura: 95-117 GB por nodo. Sugerido Fase 3: 150 GB/nodo.",
    ],
    2: [
        "Proyección BSON compacta: usar 3 KB provisionalmente.",
        "No es máximo XML: ReqPay mide 2,18 KiB realista, 5,34 KiB conservador y 9,49 KiB máximo sin firma.",
    ],
    3: [
        "Proyección compacta leída: usar 3 KB provisionalmente.",
        "RespListAccount: 2,31 KiB con 5 cuentas; supera 3 KiB desde 7 cuentas sin firma.",
    ],
    4: [
        "Supuesto preliminar: 2 horas diarias de pico.",
        "Pendiente de confirmar con la distribución horaria real.",
    ],
    5: [
        "Pico 2032: 25,75 escrituras/s conservador y 31,71 optimista.",
        "Supone 4 eventos/transacción y factor 10x. Prueba: 40 escrituras/s; estrés: 80/s.",
    ],
    6: [
        "Supuesto preliminar: 2 horas diarias, alineadas al pico transaccional.",
    ],
    7: [
        "Con 3 consultas/transacción: pico optimista 23,78 lecturas/s.",
        "Prueba recomendada: 30 lecturas/s; estrés: 60/s.",
    ],
    8: [
        "Reserva inicial: 20% de los datos.",
        "2032: 12,18 GB conservador y 15 GB optimista. Validar con collStats y $indexStats.",
    ],
    9: [
        "Hot data propuesto: 30 días.",
        "Retención total: 365 días mediante TTL.",
    ],
    10: [
        "95% consultas CRUD o búsquedas puntuales; 5% agregaciones/analítica.",
    ],
    11: [
        "Promedio: 1 documento por consulta.",
        "ListAccount puede contener varias cuentas dentro de una respuesta XML.",
    ],
    12: [
        "Supuesto del cuestionario: 20 minutos.",
        "Para el backend son más relevantes concurrencia y solicitudes por segundo.",
    ],
    13: [
        "Carga continua desde Kafka mediante upsert idempotente.",
        "Volumen anual 2027-2032: 1,1-20,3 MM conservador y 2-25 MM optimista.",
    ],
    14: [
        "Sí, propuesto y pendiente de aprobación.",
        "Proyección Kafka-MongoDB: p95 <= 500 ms; p99 <= 1 000 ms.",
    ],
    15: [
        "Lectura API p95 <= 100 ms; p99 <= 200 ms; upsert p95 <= 100 ms.",
        "Proyección extremo a extremo p95 <= 500 ms.",
    ],
    16: [
        "Estáticas y definidas en el código de la aplicación.",
        "No se prevén consultas ad-hoc de usuario.",
    ],
    17: [
        "No existe batch recurrente; la ingesta es continua.",
        "Carga inicial/backfill propuesta: 00:00-04:00, hora de Lima.",
    ],
    18: [
        "Purga automática mediante TTL a los 365 días desde completed_at.",
        "Si se requiere mayor retención, archivar antes del vencimiento.",
    ],
    19: [
        "Producción: continuous backup/PITR más snapshot diario.",
        "RPO propuesto <= 5 min; RTO <= 60 min. Pendiente de aprobación.",
    ],
    20: [
        "DEV, QA/SIT, UAT/Staging y Performance.",
        "DR se considera parte de la continuidad productiva.",
    ],
    21: [
        "DEV: 10%; QA/SIT: 25%; UAT/Staging: 50%.",
        "Performance: 100% temporal durante campañas de prueba.",
    ],
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def clone_or_none(element):
    return deepcopy(element) if element is not None else None


def replace_answer_cell(cell, lines: list[str]) -> None:
    paragraphs = cell.findall(f"{W}p")
    if not paragraphs:
        raise ValueError("Answer cell has no paragraph template")

    paragraph_templates = []
    for paragraph in paragraphs:
        p_pr = paragraph.find(f"{W}pPr")
        first_run = paragraph.find(f"{W}r")
        r_pr = first_run.find(f"{W}rPr") if first_run is not None else None
        paragraph_templates.append((clone_or_none(p_pr), clone_or_none(r_pr)))

    for paragraph in paragraphs:
        cell.remove(paragraph)

    for index, line in enumerate(lines):
        p_pr, r_pr = paragraph_templates[min(index, len(paragraph_templates) - 1)]
        paragraph = etree.Element(f"{W}p")
        if p_pr is not None:
            paragraph.append(deepcopy(p_pr))
        run = etree.SubElement(paragraph, f"{W}r")
        if r_pr is not None:
            run.append(deepcopy(r_pr))
        text = etree.SubElement(run, f"{W}t")
        if line[:1].isspace() or line[-1:].isspace():
            text.set(XML_SPACE, "preserve")
        text.text = line
        cell.append(paragraph)


def build() -> None:
    actual_hash = sha256(SOURCE)
    if actual_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            f"Reference changed: expected {EXPECTED_SOURCE_SHA256}, got {actual_hash}"
        )

    with ZipFile(SOURCE, "r") as source_zip:
        document_xml = source_zip.read("word/document.xml")
        root = etree.fromstring(document_xml)
        tables = root.xpath(".//w:tbl", namespaces=NS)
        if len(tables) != 1:
            raise ValueError(f"Expected one table, found {len(tables)}")
        rows = tables[0].xpath("./w:tr", namespaces=NS)
        if len(rows) != 22:
            raise ValueError(f"Expected 22 rows, found {len(rows)}")

        for row_index, lines in ANSWERS.items():
            cells = rows[row_index].xpath("./w:tc", namespaces=NS)
            if len(cells) != 2:
                raise ValueError(f"Row {row_index} does not have two physical cells")
            replace_answer_cell(cells[1], lines)

        updated_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(OUTPUT, "w") as output_zip:
            for info in source_zip.infolist():
                payload = updated_xml if info.filename == "word/document.xml" else source_zip.read(info.filename)
                output_zip.writestr(info, payload)

    with ZipFile(SOURCE, "r") as source_zip, ZipFile(OUTPUT, "r") as output_zip:
        source_names = source_zip.namelist()
        if source_names != output_zip.namelist():
            raise RuntimeError("Package member order or inventory changed")
        changed = []
        for name in source_names:
            if source_zip.read(name) != output_zip.read(name):
                changed.append(name)
        if changed != ["word/document.xml"]:
            raise RuntimeError(f"Unexpected changed package parts: {changed}")

    document = Document(OUTPUT)
    if len(document.sections) != 1 or len(document.tables) != 1:
        raise RuntimeError("Output structure does not match the original form")
    table = document.tables[0]
    if len(table.rows) != 22 or len(table.columns) != 2:
        raise RuntimeError("Output table is not 22x2")
    for row_index, lines in ANSWERS.items():
        expected = "\n".join(lines)
        actual = table.rows[row_index].cells[1].text
        if actual != expected:
            raise RuntimeError(f"Row {row_index} mismatch: {actual!r}")

    if sha256(SOURCE) != EXPECTED_SOURCE_SHA256:
        raise RuntimeError("Retained reference was modified")

    print(f"OUTPUT={OUTPUT}")
    print("CHANGED_PACKAGE_PARTS=word/document.xml")
    print("STRUCTURE=1 section; 1 table; 22 rows; 2 columns")
    print(f"SOURCE_SHA256={sha256(SOURCE)}")
    print(f"OUTPUT_SHA256={sha256(OUTPUT)}")


if __name__ == "__main__":
    build()
